from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from statistics import median
from typing import Any
from zipfile import BadZipFile

import cv2
import fitz
import numpy as np
from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from fastapi import HTTPException
from lxml.etree import XMLSyntaxError

from app.file_safety import MAX_FILE_SIZE, scan_file

MAX_DOCUMENT_PAGES = 60
POINTS_PER_CENTIMETER = 72 / 2.54


@dataclass(frozen=True)
class LogoCandidate:
    page: int
    bbox: list[float]
    confidence: float


@dataclass(frozen=True)
class TemplateAnalysis:
    style_spec: dict[str, Any]
    confidence: dict[str, float]
    logo_candidates: list[LogoCandidate]


@dataclass(frozen=True)
class LogoImage:
    png: bytes
    bbox: list[float]
    confidence: float


def _validate_template_file(filename: str, content: bytes) -> str:
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail={
                "code": "file_too_large",
                "message": "The template exceeds the 25 MB limit.",
                "field": "storage_path",
            },
        )
    result = scan_file(filename, content)
    if result.reasons:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "unsafe_template",
                "message": "The template is unsupported or failed file-safety checks.",
                "field": "storage_path",
            },
        )
    suffix = PurePath(filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "unsupported_template",
                "message": "Template analysis supports PDF and DOCX files.",
                "field": "storage_path",
            },
        )
    return suffix


def _round(value: float, digits: int = 2) -> float:
    return round(float(value), digits)


def _pdf_styles(document: fitz.Document) -> tuple[dict[str, Any], dict[str, float]]:
    font_counts: Counter[tuple[str, float]] = Counter()
    span_rows: list[tuple[str, float, bool, str, tuple[float, ...], float]] = []
    page_margins: list[list[float]] = []
    headers: list[str] = []
    footers: list[str] = []

    for page in document:
        page_dict = page.get_text("dict")
        page_spans: list[dict[str, Any]] = []
        for block in page_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                line_height = float(line["bbox"][3] - line["bbox"][1])
                for span in line.get("spans", []):
                    text = span.get("text", "").strip()
                    if not text:
                        continue
                    font = str(span.get("font", ""))
                    size = float(span.get("size", 0))
                    bold = "bold" in font.lower() or "black" in font.lower()
                    bbox = tuple(float(value) for value in span["bbox"])
                    page_spans.append(span)
                    span_rows.append((font, size, bold, text, bbox, line_height))
                    font_counts[(font, round(size, 1))] += max(1, len(text))
                    if bbox[1] < page.rect.height * 0.12:
                        headers.append(text)
                    if bbox[3] > page.rect.height * 0.88:
                        footers.append(text)

        if page_spans:
            boxes = [
                tuple(float(value) for value in span["bbox"])
                for span in page_spans
            ]
            page_margins.append(
                [
                    min(box[1] for box in boxes),
                    page.rect.width - max(box[2] for box in boxes),
                    page.rect.height - max(box[3] for box in boxes),
                    min(box[0] for box in boxes),
                ]
            )

    if not span_rows:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "template_has_no_text",
                "message": "No extractable text was found in the PDF.",
                "field": "storage_path",
            },
        )

    body_font, body_size = font_counts.most_common(1)[0][0]
    body_rows = [row for row in span_rows if row[0] == body_font and row[1] == body_size]
    heading_rows = [
        row for row in span_rows if row[2] or row[1] >= body_size + 1
    ]
    if not heading_rows:
        largest_size = max(row[1] for row in span_rows)
        heading_rows = [row for row in span_rows if row[1] == largest_size]
    heading_counts = Counter(
        (row[0], round(row[1], 1), row[2]) for row in heading_rows
    )
    heading_font, heading_size, heading_bold = heading_counts.most_common(1)[0][0]
    line_ratios = [
        row[5] / row[1]
        for row in body_rows
        if row[1] > 0 and row[5] > 0
    ]
    margins = (
        [
            _round(median(values) / POINTS_PER_CENTIMETER)
            for values in zip(*page_margins)
        ]
        if page_margins
        else []
    )

    style_spec = {
        "font": body_font,
        "size_pt": body_size,
        "heading": {
            "font": heading_font,
            "size_pt": heading_size,
            "bold": heading_bold,
        },
        "line_spacing": _round(median(line_ratios)) if line_ratios else None,
        "margins_cm": margins,
        "header": " ".join(dict.fromkeys(headers)) or None,
        "footer": " ".join(dict.fromkeys(footers)) or None,
        "logo_path": None,
        "columns": [],
    }
    body_style_confidence = _round(
        font_counts[(body_font, body_size)] / sum(font_counts.values()), 2
    )
    confidence = {
        "font": body_style_confidence,
        "size_pt": body_style_confidence,
        "heading": _round(len(heading_rows) / len(span_rows), 2),
        "line_spacing": 0.8 if line_ratios else 0.0,
        "margins_cm": _round(len(page_margins) / len(document), 2),
        "header": 0.8 if headers else 0.0,
        "footer": 0.8 if footers else 0.0,
        "columns": 0.0,
    }
    return style_spec, confidence


def _docx_styles(content: bytes) -> tuple[dict[str, Any], dict[str, float]]:
    try:
        document = WordDocument(BytesIO(content))
    except (
        AttributeError,
        BadZipFile,
        PackageNotFoundError,
        ValueError,
        OSError,
        XMLSyntaxError,
    ) as error:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "invalid_template",
                "message": "The DOCX template could not be opened.",
                "field": "storage_path",
            },
        ) from error

    normal = document.styles["Normal"]
    body_font = normal.font.name
    body_size = normal.font.size.pt if normal.font.size else None
    heading_style = document.styles["Heading 1"]
    heading_font = heading_style.font.name or body_font
    heading_size = heading_style.font.size.pt if heading_style.font.size else None
    section = document.sections[0] if document.sections else None
    margins = (
        [
            _round(section.top_margin.cm),
            _round(section.right_margin.cm),
            _round(section.bottom_margin.cm),
            _round(section.left_margin.cm),
        ]
        if section
        else []
    )

    line_spacing_values = []
    for paragraph in document.paragraphs:
        line_spacing = paragraph.paragraph_format.line_spacing
        if isinstance(line_spacing, float):
            line_spacing_values.append(line_spacing)
        elif line_spacing is not None and body_size:
            line_spacing_values.append(line_spacing.pt / body_size)
    header = (
        " ".join(
            paragraph.text.strip()
            for paragraph in section.header.paragraphs
            if paragraph.text.strip()
        )
        if section
        else ""
    )
    footer = (
        " ".join(
            paragraph.text.strip()
            for paragraph in section.footer.paragraphs
            if paragraph.text.strip()
        )
        if section
        else ""
    )
    columns = []
    if document.tables:
        first_row = document.tables[0].rows[0]
        columns = [
            cell.text.strip() or f"column_{index + 1}"
            for index, cell in enumerate(first_row.cells)
        ]

    style_spec = {
        "font": body_font,
        "size_pt": _round(body_size) if body_size else None,
        "heading": {
            "font": heading_font,
            "size_pt": _round(heading_size) if heading_size else None,
            "bold": bool(heading_style.font.bold),
        },
        "line_spacing": (
            _round(median(line_spacing_values)) if line_spacing_values else None
        ),
        "margins_cm": margins,
        "header": header or None,
        "footer": footer or None,
        "logo_path": None,
        "columns": columns,
    }
    confidence = {
        "font": 0.95 if body_font else 0.0,
        "size_pt": 0.95 if body_size else 0.0,
        "heading": 0.85 if heading_font or heading_size else 0.0,
        "line_spacing": 0.8 if line_spacing_values else 0.0,
        "margins_cm": 0.95 if margins else 0.0,
        "header": 0.9 if header else 0.0,
        "footer": 0.9 if footer else 0.0,
        "columns": 0.75 if columns else 0.0,
    }
    return style_spec, confidence


def _box_iou(left: list[float], right: list[float]) -> float:
    x0 = max(left[0], right[0])
    y0 = max(left[1], right[1])
    x1 = min(left[2], right[2])
    y1 = min(left[3], right[3])
    intersection = max(0.0, x1 - x0) * max(0.0, y1 - y0)
    left_area = max(0.0, left[2] - left[0]) * max(0.0, left[3] - left[1])
    right_area = max(0.0, right[2] - right[0]) * max(0.0, right[3] - right[1])
    union = left_area + right_area - intersection
    return intersection / union if union else 0.0


def _detect_page_logos(page: fitz.Page, page_number: int) -> list[LogoCandidate]:
    page_rect = page.rect
    header_rect = fitz.Rect(
        page_rect.x0,
        page_rect.y0,
        page_rect.x1,
        page_rect.y0 + page_rect.height * 0.35,
    )
    candidates: list[LogoCandidate] = []

    for image in page.get_images(full=True):
        for image_rect in page.get_image_rects(image[0]):
            rect = image_rect & header_rect
            area_ratio = (rect.width * rect.height) / (page_rect.width * page_rect.height)
            if (
                rect.is_empty
                or rect.width > page_rect.width * 0.7
                or not 0.001 <= area_ratio <= 0.15
            ):
                continue
            candidates.append(
                LogoCandidate(
                    page=page_number,
                    bbox=[
                        _round(rect.x0),
                        _round(rect.y0),
                        _round(rect.x1),
                        _round(rect.y1),
                    ],
                    confidence=0.92,
                )
            )

    pixmap = page.get_pixmap(
        matrix=fitz.Matrix(2, 2),
        clip=header_rect,
        colorspace=fitz.csRGB,
        alpha=False,
    )
    pixels = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
        pixmap.height, pixmap.width, pixmap.n
    )
    gray = cv2.cvtColor(pixels, cv2.COLOR_RGB2GRAY)
    _, binary = cv2.threshold(
        gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU
    )
    kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (max(5, pixmap.width // 160), max(3, pixmap.height // 90)),
    )
    merged = cv2.morphologyEx(binary, cv2.MORPH_CLOSE, kernel)
    contours, _ = cv2.findContours(
        merged, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
    )
    for contour in contours:
        x, y, width, height = cv2.boundingRect(contour)
        area_ratio = (width * height) / (pixmap.width * pixmap.height)
        aspect_ratio = width / height if height else 0
        if not 0.002 <= area_ratio <= 0.12 or not 0.35 <= aspect_ratio <= 5.0:
            continue
        box = [
            _round(header_rect.x0 + x / 2),
            _round(header_rect.y0 + y / 2),
            _round(header_rect.x0 + (x + width) / 2),
            _round(header_rect.y0 + (y + height) / 2),
        ]
        if any(_box_iou(box, candidate.bbox) > 0.6 for candidate in candidates):
            continue
        area_confidence = min(area_ratio / 0.02, 1.0)
        candidates.append(
            LogoCandidate(
                page=page_number,
                bbox=box,
                confidence=_round(0.55 + 0.25 * area_confidence, 2),
            )
        )

    return sorted(candidates, key=lambda item: item.confidence, reverse=True)[:5]


def analyse_template(filename: str, content: bytes) -> TemplateAnalysis:
    suffix = _validate_template_file(filename, content)
    if suffix == ".docx":
        style_spec, confidence = _docx_styles(content)
        return TemplateAnalysis(style_spec, confidence, [])

    try:
        document = fitz.open(stream=content, filetype="pdf")
    except (fitz.FileDataError, RuntimeError, ValueError) as error:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "invalid_template",
                "message": "The PDF template could not be opened.",
                "field": "storage_path",
            },
        ) from error
    with document:
        if document.needs_pass:
            raise HTTPException(
                status_code=422,
                detail={
                    "code": "encrypted_template",
                    "message": "Password-protected PDFs cannot be analysed.",
                    "field": "storage_path",
                },
            )
        if not 1 <= document.page_count <= MAX_DOCUMENT_PAGES:
            raise HTTPException(
                status_code=422,
                detail={
                    "code": "invalid_page_count",
                    "message": "PDF templates must contain between 1 and 60 pages.",
                    "field": "storage_path",
                },
            )
        style_spec, confidence = _pdf_styles(document)
        candidates = [
            candidate
            for index in range(min(document.page_count, 10))
            for candidate in _detect_page_logos(document[index], index + 1)
        ]
        candidates.sort(key=lambda item: item.confidence, reverse=True)
        return TemplateAnalysis(style_spec, confidence, candidates[:20])


def extract_logo(
    filename: str,
    content: bytes,
    page_number: int,
    crop: list[float] | None = None,
) -> LogoImage:
    suffix = _validate_template_file(filename, content)
    if suffix != ".pdf":
        raise HTTPException(
            status_code=422,
            detail={
                "code": "unsupported_template",
                "message": "Logo extraction currently supports PDF templates.",
                "field": "storage_path",
            },
        )
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except (fitz.FileDataError, RuntimeError, ValueError) as error:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "invalid_template",
                "message": "The PDF template could not be opened.",
                "field": "storage_path",
            },
        ) from error

    with document:
        if document.needs_pass:
            raise HTTPException(
                status_code=422,
                detail={
                    "code": "encrypted_template",
                    "message": "Password-protected PDFs cannot be analysed.",
                    "field": "storage_path",
                },
            )
        if not 1 <= page_number <= min(document.page_count, MAX_DOCUMENT_PAGES):
            raise HTTPException(
                status_code=422,
                detail={
                    "code": "invalid_page",
                    "message": "The requested page must be between 1 and 60.",
                    "field": "page",
                },
            )

        page = document[page_number - 1]
        if crop is not None:
            if len(crop) != 4 or not all(np.isfinite(value) for value in crop):
                raise HTTPException(
                    status_code=422,
                    detail={
                        "code": "invalid_crop",
                        "message": "Crop must contain four finite page coordinates.",
                        "field": "crop",
                    },
                )
            rect = fitz.Rect(*crop) & page.rect
            if rect.is_empty:
                raise HTTPException(
                    status_code=422,
                    detail={
                        "code": "invalid_crop",
                        "message": "The crop must overlap the page.",
                        "field": "crop",
                    },
                )
            confidence = 0.95
        else:
            candidates = _detect_page_logos(page, page_number)
            if not candidates:
                raise HTTPException(
                    status_code=404,
                    detail={
                        "code": "logo_not_found",
                        "message": "No logo candidate was found on the requested page.",
                        "field": "page",
                    },
                )
            candidate = candidates[0]
            rect = fitz.Rect(*candidate.bbox)
            confidence = candidate.confidence

        pixmap = page.get_pixmap(
            matrix=fitz.Matrix(3, 3),
            clip=rect,
            colorspace=fitz.csRGB,
            alpha=False,
        )
        pixels = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
            pixmap.height, pixmap.width, pixmap.n
        )
        bgr = cv2.cvtColor(pixels, cv2.COLOR_RGB2BGR)
        encoded, output = cv2.imencode(".png", bgr)
        if not encoded:
            raise HTTPException(
                status_code=500,
                detail={
                    "code": "logo_encoding_failed",
                    "message": "The detected logo could not be encoded as PNG.",
                },
            )
        return LogoImage(
            png=output.tobytes(),
            bbox=[_round(rect.x0), _round(rect.y0), _round(rect.x1), _round(rect.y1)],
            confidence=confidence,
        )
