import json
import logging
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any
from zipfile import ZipFile

import fitz
import httpx
import jwt
import numpy as np
import pytest
import cv2
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app import storage
from app.main import app
from app.template_analysis import MAX_DOCUMENT_PAGES, analyse_template, extract_logo

client = TestClient(app)
JWT_SECRET = "template-test-secret-with-at-least-thirty-two-bytes"
SUPABASE_URL = "https://storage.test"
SUPABASE_ANON_KEY = "test-anon-key"


def auth_headers() -> dict[str, str]:
    now = datetime.now(timezone.utc)
    token = jwt.encode(
        {
            "sub": "template-test-user",
            "role": "authenticated",
            "aud": "authenticated",
            "iss": f"{SUPABASE_URL}/auth/v1",
            "iat": now,
            "exp": now + timedelta(minutes=5),
        },
        JWT_SECRET,
        algorithm="HS256",
    )
    return {"Authorization": f"Bearer {token}"}


def make_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    page.draw_rect(fitz.Rect(48, 40, 150, 96), color=(0.05, 0.25, 0.5), fill=(0.05, 0.25, 0.5))
    page.insert_text((56, 74), "NUST", fontsize=24, fontname="hebo", color=(1, 1, 1))
    page.insert_text((50, 150), "Research Template", fontsize=18, fontname="hebo")
    page.insert_text((50, 205), "Study information and participant procedures.", fontsize=11)
    page.insert_text((50, 750), "Ethics reference: TEST-001", fontsize=9)
    output = document.tobytes()
    document.close()
    return output


def make_pdf_pages(page_count: int, *, include_text: bool = True) -> bytes:
    document = fitz.open()
    image, encoded_image = cv2.imencode(
        ".png",
        np.full((32, 32, 3), 255, dtype=np.uint8),
    )
    assert image
    for page_number in range(page_count):
        page = document.new_page(width=612, height=792)
        if include_text:
            page.insert_text((50, 150), f"Template page {page_number + 1}")
        else:
            page.insert_image(
                fitz.Rect(40, 40, 300, 300),
                stream=encoded_image.tobytes(),
            )
    output = document.tobytes()
    document.close()
    return output


def make_zero_page_pdf() -> bytes:
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in [
        (1, b"<< /Type /Catalog /Pages 2 0 R >>"),
        (2, b"<< /Type /Pages /Kids [] /Count 0 >>"),
    ]:
        offsets.append(len(content))
        content.extend(f"{object_number} 0 obj\n".encode())
        content.extend(body)
        content.extend(b"\nendobj\n")

    xref_offset = len(content)
    content.extend(b"xref\n0 3\n0000000000 65535 f \n")
    content.extend(
        b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    )
    content.extend(
        f"trailer\n<< /Size 3 /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(content)


def make_encrypted_pdf() -> bytes:
    document = fitz.open()
    document.new_page()
    output = document.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="test-owner",
        user_pw="test-user",
    )
    document.close()
    return output


def make_docx() -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = 11 * 12700
    document.styles["Heading 1"].font.name = "Arial"
    document.styles["Heading 1"].font.size = 14 * 12700
    document.add_heading("Research Template", level=1)
    document.add_paragraph("Study information and participant procedures.")
    document.sections[0].header.paragraphs[0].text = "Test Institution"
    document.sections[0].footer.paragraphs[0].text = "Ethics reference"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_docx_without_part(part_name: str) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(make_docx())) as source, ZipFile(output, "w") as target:
        for item in source.infolist():
            if item.filename != part_name:
                target.writestr(item, source.read(item.filename))
    return output.getvalue()


@pytest.fixture(autouse=True)
def configure_test_auth(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("SUPABASE_JWT_SECRET", JWT_SECRET)
    monkeypatch.delenv("SUPABASE_URL", raising=False)
    monkeypatch.setenv("SUPABASE_URL", SUPABASE_URL)
    monkeypatch.setenv("SUPABASE_ANON_KEY", SUPABASE_ANON_KEY)
    monkeypatch.setenv("SUPABASE_TEMPLATE_BUCKET", "templates")


def mock_storage_api(
    monkeypatch: pytest.MonkeyPatch,
    handler: Any,
) -> None:
    async_client = httpx.AsyncClient
    transport = httpx.MockTransport(handler)

    def make_client(**kwargs: Any) -> httpx.AsyncClient:
        return async_client(transport=transport, **kwargs)

    monkeypatch.setattr(storage.httpx, "AsyncClient", make_client)


def test_parse_storage_path_requires_bucket_and_object() -> None:
    parsed = storage.parse_storage_path("templates/research/sample.pdf")

    assert parsed.bucket == "templates"
    assert parsed.object_path == "research/sample.pdf"
    assert parsed.name == "sample.pdf"


@pytest.mark.parametrize(
    "storage_path",
    ["", "templates", "/templates/sample.pdf", "templates/../private.pdf", "templates\\sample.pdf"],
)
def test_parse_storage_path_rejects_invalid_paths(storage_path: str) -> None:
    with pytest.raises(HTTPException) as error:
        storage.parse_storage_path(storage_path)

    assert error.value.status_code == 422


def test_analyse_pdf_extracts_styles_and_logo_candidate() -> None:
    analysis = analyse_template("sample.pdf", make_pdf())

    assert analysis.style_spec["font"]
    assert analysis.style_spec["size_pt"] == 11
    assert analysis.style_spec["heading"]["size_pt"] >= 18
    assert analysis.style_spec["footer"] == "Ethics reference: TEST-001"
    assert analysis.confidence["font"] > 0
    assert analysis.logo_candidates
    assert analysis.logo_candidates[0].page == 1
    assert len(analysis.logo_candidates[0].bbox) == 4


def test_analyse_docx_extracts_styles_header_footer_and_margins() -> None:
    analysis = analyse_template("sample.docx", make_docx())

    assert analysis.style_spec["font"] == "Arial"
    assert analysis.style_spec["size_pt"] == 11
    assert analysis.style_spec["heading"]["size_pt"] == 14
    assert analysis.style_spec["margins_cm"]
    assert analysis.style_spec["header"] == "Test Institution"
    assert analysis.style_spec["footer"] == "Ethics reference"
    assert analysis.logo_candidates == []


def test_analyse_rejects_malformed_docx() -> None:
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "not XML")

    with pytest.raises(HTTPException) as error:
        analyse_template("sample.docx", output.getvalue())

    assert error.value.status_code == 422
    assert error.value.detail["code"] == "invalid_template"


def test_analyse_rejects_docx_missing_required_styles_part() -> None:
    with pytest.raises(HTTPException) as error:
        analyse_template("missing-styles.docx", make_docx_without_part("word/styles.xml"))

    assert error.value.status_code == 422
    assert error.value.detail == {
        "code": "invalid_template",
        "message": "The DOCX template could not be opened.",
        "field": "storage_path",
    }


@pytest.mark.parametrize(
    ("filename", "content", "error_code"),
    [
        ("encrypted.pdf", make_encrypted_pdf(), "encrypted_template"),
        ("corrupt.pdf", b"%PDF-1.7\nthis is not a valid PDF", "invalid_template"),
        (
            "scanned.pdf",
            make_pdf_pages(1, include_text=False),
            "template_has_no_text",
        ),
    ],
    ids=["encrypted-pdf", "corrupt-pdf", "image-only-pdf"],
)
def test_analyse_returns_documented_errors_for_pdf_edge_cases(
    filename: str,
    content: bytes,
    error_code: str,
) -> None:
    with pytest.raises(HTTPException) as error:
        analyse_template(filename, content)

    assert error.value.status_code == 422
    assert error.value.detail["code"] == error_code
    assert error.value.detail["field"] == "storage_path"


def test_analyse_accepts_pdf_at_page_limit() -> None:
    analysis = analyse_template(
        "limit.pdf", make_pdf_pages(MAX_DOCUMENT_PAGES)
    )

    assert analysis.style_spec["font"] is not None


def test_analyse_rejects_pdf_over_page_limit() -> None:
    with pytest.raises(HTTPException) as error:
        analyse_template(
            "over-limit.pdf", make_pdf_pages(MAX_DOCUMENT_PAGES + 1)
        )

    assert error.value.status_code == 422
    assert error.value.detail == {
        "code": "invalid_page_count",
        "message": "PDF templates must contain between 1 and 60 pages.",
        "field": "storage_path",
    }


def test_analyse_rejects_zero_page_pdf() -> None:
    with pytest.raises(HTTPException) as error:
        analyse_template("zero-pages.pdf", make_zero_page_pdf())

    assert error.value.status_code == 422
    assert error.value.detail["code"] == "invalid_page_count"


@pytest.mark.parametrize(
    ("filename", "signature"),
    [("oversized.pdf", b"%PDF-"), ("oversized.docx", b"PK")],
    ids=["pdf", "docx"],
)
def test_analyse_rejects_document_over_size_limit(
    filename: str, signature: bytes
) -> None:
    with pytest.raises(HTTPException) as error:
        analyse_template(filename, signature + b"x" * (25 * 1024 * 1024))

    assert error.value.status_code == 413
    assert error.value.detail["code"] == "file_too_large"


def test_extract_logo_rejects_pdf_over_page_limit() -> None:
    with pytest.raises(HTTPException) as error:
        extract_logo(
            "over-limit.pdf",
            make_pdf_pages(MAX_DOCUMENT_PAGES + 1),
            page_number=1,
            crop=[0, 0, 100, 100],
        )

    assert error.value.status_code == 422
    assert error.value.detail["code"] == "invalid_page_count"


@pytest.mark.parametrize(
    ("filename", "content", "error_code"),
    [
        ("encrypted.pdf", make_encrypted_pdf(), "encrypted_template"),
        ("corrupt.pdf", b"%PDF-1.7\nthis is not a valid PDF", "invalid_template"),
        (
            "scanned.pdf",
            make_pdf_pages(1, include_text=False),
            "template_has_no_text",
        ),
        (
            "over-limit.pdf",
            make_pdf_pages(MAX_DOCUMENT_PAGES + 1),
            "invalid_page_count",
        ),
        (
            "missing-styles.docx",
            make_docx_without_part("word/styles.xml"),
            "invalid_template",
        ),
    ],
    ids=["encrypted-pdf", "corrupt-pdf", "image-only-pdf", "too-many-pages", "malformed-docx"],
)
def test_upload_endpoint_returns_edge_case_errors_without_storage_write(
    monkeypatch: pytest.MonkeyPatch,
    filename: str,
    content: bytes,
    error_code: str,
) -> None:
    def fail_if_storage_called(*args: Any, **kwargs: Any) -> str:
        pytest.fail("Invalid templates must not be sent to Storage.")

    monkeypatch.setattr(storage, "upload_storage_object", fail_if_storage_called)
    response = client.post(
        "/v1/templates/upload",
        files={"file": (filename, content)},
        headers=auth_headers(),
    )

    assert response.status_code == 422
    assert response.json()["code"] == error_code
    assert response.json()["field"] == "storage_path"


def test_extract_logo_returns_png_with_requested_bbox() -> None:
    logo = extract_logo("sample.pdf", make_pdf(), page_number=1, crop=[48, 40, 150, 96])

    assert logo.png.startswith(b"\x89PNG\r\n\x1a\n")
    assert logo.bbox == [48.0, 40.0, 150.0, 96.0]
    assert logo.confidence == 0.95


@pytest.mark.parametrize(
    ("page", "crop"),
    [(0, None), (61, None), (1, [1000, 1000, 1100, 1100]), (1, [1, 2, 3])],
)
def test_extract_logo_rejects_invalid_page_or_crop(
    page: int, crop: list[float] | None
) -> None:
    with pytest.raises(HTTPException) as error:
        extract_logo("sample.pdf", make_pdf(), page, crop)

    assert error.value.status_code == 422


def test_analyse_endpoint_uses_caller_jwt_for_storage(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    headers = auth_headers()
    expected_token = headers["Authorization"]

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.method == "GET"
        assert request.url.path == (
            "/storage/v1/object/authenticated/templates/research/sample.pdf"
        )
        assert request.headers["authorization"] == expected_token
        assert request.headers["apikey"] == SUPABASE_ANON_KEY
        return httpx.Response(200, content=make_pdf(), request=request)

    mock_storage_api(monkeypatch, handler)
    response = client.post(
        "/v1/templates/analyse",
        json={"storage_path": "templates/research/sample.pdf"},
        headers=headers,
    )

    assert response.status_code == 200
    assert response.json()["font"]
    assert response.json()["confidence"]["font"] > 0
    assert response.json()["logo_candidates"]


@pytest.mark.parametrize(
    ("filename", "contents", "detected_type"),
    [
        ("sample.pdf", make_pdf(), "application/pdf"),
        (
            "sample.docx",
            make_docx(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
    ],
    ids=["pdf", "docx"],
)
def test_upload_endpoint_scans_analyses_and_stores_with_caller_jwt(
    monkeypatch: pytest.MonkeyPatch,
    filename: str,
    contents: bytes,
    detected_type: str,
) -> None:
    headers = auth_headers()
    expected_token = headers["Authorization"]
    uploaded: dict[str, Any] = {}

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.method == "POST"
        assert request.headers["authorization"] == expected_token
        assert request.headers["apikey"] == SUPABASE_ANON_KEY
        assert request.headers["content-type"] == detected_type
        uploaded.update(path=request.url.path, content=request.content)
        return httpx.Response(200, json={"Key": request.url.path}, request=request)

    mock_storage_api(monkeypatch, handler)
    response = client.post(
        "/v1/templates/upload",
        files={"file": (filename, contents)},
        headers=headers,
    )

    assert response.status_code == 200
    result = response.json()
    assert result["detected_type"] == detected_type
    assert result["storage_path"].startswith("templates/uploads/")
    assert result["font"]
    assert "confidence" in result
    assert uploaded["path"].startswith("/storage/v1/object/templates/uploads/")
    assert uploaded["content"] == contents


def test_storage_failure_log_excludes_request_and_upstream_data(
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    headers = auth_headers()

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            500,
            text="sensitive upstream response",
            request=request,
        )

    mock_storage_api(monkeypatch, handler)
    caplog.set_level(logging.ERROR, logger="consentlink.service")
    response = client.post(
        "/v1/templates/upload",
        files={"file": ("private-name.pdf", make_pdf())},
        headers=headers,
    )

    assert response.status_code == 500
    records = [
        json.loads(record.message)
        for record in caplog.records
        if record.name == "consentlink.service"
    ]
    assert records == [
        {
            "error_type": "HTTPStatusError",
            "event": "storage_request_failed",
            "operation": "upload",
            "upstream_status": 500,
        },
        {
            "error_type": "HTTPException",
            "event": "http_request_failed",
            "method": "POST",
            "route": "/v1/templates/upload",
            "status_code": 500,
        },
    ]
    logged_output = caplog.text
    assert "private-name.pdf" not in logged_output
    assert "sensitive upstream response" not in logged_output
    assert headers["Authorization"] not in logged_output


def test_upload_endpoint_rejects_unsafe_file_without_storage_write(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_if_storage_called(*args: Any, **kwargs: Any) -> str:
        pytest.fail("Unsafe uploads must not be sent to Storage.")

    monkeypatch.setattr(storage, "upload_storage_object", fail_if_storage_called)
    response = client.post(
        "/v1/templates/upload",
        files={"file": ("unsafe.exe", b"MZ\x90\x00")},
        headers=auth_headers(),
    )

    assert response.status_code == 422
    assert response.json()["code"] == "unsafe_template"
    assert "executable" in response.json()["reasons"]


def test_upload_endpoint_rejects_non_template_without_storage_write(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_if_storage_called(*args: Any, **kwargs: Any) -> str:
        pytest.fail("Unsupported templates must not be sent to Storage.")

    monkeypatch.setattr(storage, "upload_storage_object", fail_if_storage_called)
    response = client.post(
        "/v1/templates/upload",
        files={"file": ("scan.png", b"\x89PNG\r\n\x1a\nimage bytes")},
        headers=auth_headers(),
    )

    assert response.status_code == 422
    assert response.json()["code"] == "unsupported_template"


def test_upload_endpoint_rejects_oversized_file_without_storage_write(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_if_storage_called(*args: Any, **kwargs: Any) -> str:
        pytest.fail("Oversized templates must not be sent to Storage.")

    monkeypatch.setattr(storage, "upload_storage_object", fail_if_storage_called)
    response = client.post(
        "/v1/templates/upload",
        files={"file": ("sample.pdf", b"%PDF-" + b"x" * (25 * 1024 * 1024))},
        headers=auth_headers(),
    )

    assert response.status_code == 413
    assert response.json()["code"] == "file_too_large"


def test_upload_endpoint_requires_authentication() -> None:
    response = client.post(
        "/v1/templates/upload",
        files={"file": ("sample.pdf", make_pdf())},
    )

    assert response.status_code == 401
    assert response.json()["code"] == "unauthorized"


def test_upload_endpoint_requires_template_bucket_configuration(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("SUPABASE_TEMPLATE_BUCKET")
    response = client.post(
        "/v1/templates/upload",
        files={"file": ("sample.pdf", make_pdf())},
        headers=auth_headers(),
    )

    assert response.status_code == 503
    assert response.json()["code"] == "template_storage_not_configured"


def test_upload_endpoint_reports_storage_permission_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(403, request=request)

    mock_storage_api(monkeypatch, handler)
    response = client.post(
        "/v1/templates/upload",
        files={"file": ("sample.pdf", make_pdf())},
        headers=auth_headers(),
    )

    assert response.status_code == 403
    assert response.json() == {
        "code": "storage_forbidden",
        "message": "The caller cannot access this storage object.",
    }


def test_analyse_endpoint_requires_authentication() -> None:
    response = client.post(
        "/v1/templates/analyse",
        json={"storage_path": "templates/sample.pdf"},
    )

    assert response.status_code == 401
    assert response.json()["code"] == "unauthorized"


def test_extract_endpoint_writes_png_and_returns_storage_path(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    headers = auth_headers()
    expected_token = headers["Authorization"]
    uploaded: dict[str, Any] = {}

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["authorization"] == expected_token
        assert request.headers["apikey"] == SUPABASE_ANON_KEY
        if request.method == "GET":
            assert request.url.path == (
                "/storage/v1/object/authenticated/templates/research/sample.pdf"
            )
            return httpx.Response(200, content=make_pdf(), request=request)
        uploaded.update(path=request.url.path, content=request.content)
        return httpx.Response(200, json={"Key": "templates/logos/result.png"}, request=request)

    mock_storage_api(monkeypatch, handler)
    response = client.post(
        "/v1/logos/extract",
        json={
            "storage_path": "templates/research/sample.pdf",
            "page": 1,
            "crop": [48, 40, 150, 96],
        },
        headers=headers,
    )

    assert response.status_code == 200
    assert response.json()["png_path"].startswith("templates/logos/")
    assert response.json()["bbox"] == [48.0, 40.0, 150.0, 96.0]
    assert uploaded["path"].startswith("/storage/v1/object/templates/logos/")
    assert uploaded["content"].startswith(b"\x89PNG\r\n\x1a\n")


def test_uploaded_template_path_can_be_used_for_logo_extraction(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    headers = auth_headers()
    expected_token = headers["Authorization"]
    stored_objects: dict[str, bytes] = {}
    upload_calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal upload_calls
        assert request.headers["authorization"] == expected_token
        assert request.headers["apikey"] == SUPABASE_ANON_KEY
        if request.method == "GET":
            object_path = request.url.path.removeprefix(
                "/storage/v1/object/authenticated/"
            )
            if object_path not in stored_objects:
                return httpx.Response(404, request=request)
            return httpx.Response(200, content=stored_objects[object_path], request=request)
        object_path = request.url.path.removeprefix("/storage/v1/object/")
        stored_objects[object_path] = request.content
        upload_calls += 1
        return httpx.Response(200, json={"Key": object_path}, request=request)

    mock_storage_api(monkeypatch, handler)
    upload_response = client.post(
        "/v1/templates/upload",
        files={"file": ("sample.pdf", make_pdf())},
        headers=headers,
    )
    assert upload_response.status_code == 200
    source_path = upload_response.json()["storage_path"]

    logo_response = client.post(
        "/v1/logos/extract",
        json={"storage_path": source_path, "page": 1, "crop": [48, 40, 150, 96]},
        headers=headers,
    )

    assert logo_response.status_code == 200
    logo_path = logo_response.json()["png_path"]
    assert logo_path.startswith("templates/logos/")
    assert stored_objects[logo_path].startswith(b"\x89PNG\r\n\x1a\n")
    assert upload_calls == 2


def test_analyse_endpoint_rejects_missing_storage_path() -> None:
    response = client.post(
        "/v1/templates/analyse",
        json={},
        headers=auth_headers(),
    )

    assert response.status_code == 422
    assert response.json()["code"] == "invalid_request"
    assert response.json()["field"] == "storage_path"
